"""The few OOXML constructs python-docx exposes no API for.

These reach into ``paragraph._p``, ``table._tbl`` and ``style.element`` on
purpose -- there is no public route to hyperlinks, borders, cell margins or run
language. Keeping them in one module means the private-API surface is a handful
of functions wide instead of scattered through the renderer.

Element *order* matters: WordprocessingML declares a fixed child sequence, and
Word silently "repairs" documents that violate it -- damage nobody notices until
a recruiter opens the file. python-docx orders the elements it knows about, so
these helpers only have to place their own.
"""

from __future__ import annotations

from typing import Literal

from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.shared import Length
from docx.styles.style import BaseStyle
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run

# https://docs.microsoft.com/openspecs/office_standards/ms-oi29500 -- CT_PPr
# children that must precede w:pBdr.
_BEFORE_PBDR = frozenset(
    qn(tag)
    for tag in (
        "w:pStyle",
        "w:keepNext",
        "w:keepLines",
        "w:pageBreakBefore",
        "w:framePr",
        "w:widowControl",
        "w:numPr",
        "w:suppressLineNumbers",
    )
)

# CT_TblPr children that must precede w:tblBorders, and then the further ones
# that must also precede w:tblCellMar.
_BEFORE_TBL_BORDERS = frozenset(
    qn(tag)
    for tag in (
        "w:tblStyle",
        "w:tblpPr",
        "w:tblOverlap",
        "w:bidiVisual",
        "w:tblStyleRowBandSize",
        "w:tblStyleColBandSize",
        "w:tblW",
        "w:jc",
        "w:tblCellSpacing",
        "w:tblInd",
    )
)
_BEFORE_TBL_CELL_MAR = _BEFORE_TBL_BORDERS | {
    qn(tag) for tag in ("w:tblBorders", "w:shd", "w:tblLayout")
}

HYPERLINK_RELATIONSHIP = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"
)

TWIPS_PER_POINT = 20


def _insert_after(
    parent: BaseOxmlElement, element: BaseOxmlElement, precede: frozenset[str]
) -> None:
    """Insert ``element`` after the last child of ``parent`` that must precede it."""
    index = 0
    for position, child in enumerate(parent):
        if child.tag in precede:
            index = position + 1
    parent.insert(index, element)


def _single_bottom(*, color: str, size: int, space: int) -> BaseOxmlElement:
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    return bottom


def add_bottom_border(paragraph: Paragraph, *, color: str, size: int = 6, space: int = 2) -> None:
    """Draw a rule under ``paragraph``.

    Args:
        color: RGB hex without a leading ``#``.
        size: Line width in eighths of a point, so ``12`` is 1.5pt.
        space: Gap between text and rule, in points.
    """
    borders = OxmlElement("w:pBdr")
    borders.append(_single_bottom(color=color, size=size, space=space))
    _insert_after(paragraph._p.get_or_add_pPr(), borders, _BEFORE_PBDR)


def add_table_bottom_border(table: Table, *, color: str, size: int = 6) -> None:
    """Draw a rule under the whole of ``table``, spanning every column.

    The paragraph-level equivalent would only underline one cell, so a header
    laid out as columns needs the rule on the table instead.
    """
    borders = OxmlElement("w:tblBorders")
    borders.append(_single_bottom(color=color, size=size, space=0))
    _insert_after(table._tbl.tblPr, borders, _BEFORE_TBL_BORDERS)


def set_table_cell_margins(
    table: Table,
    *,
    top: float = 0.0,
    left: float = 0.0,
    bottom: float = 0.0,
    right: float = 0.0,
) -> None:
    """Set ``table``'s cell insets, in points.

    Word's default left inset would indent a layout table's text relative to the
    paragraphs around it; zeroing it keeps a two-column header aligned with the
    single-column body.

    The edges are keyword-only and written in the order ``CT_TblCellMar``
    declares -- top, left, bottom, right -- because that sequence is part of the
    schema too, not just the position of ``w:tblCellMar`` among its siblings.
    """
    element = OxmlElement("w:tblCellMar")
    for edge, points in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        child = OxmlElement(f"w:{edge}")
        child.set(qn("w:w"), str(round(points * TWIPS_PER_POINT)))
        child.set(qn("w:type"), "dxa")
        element.append(child)
    _insert_after(table._tbl.tblPr, element, _BEFORE_TBL_CELL_MAR)


def set_column_widths(table: Table, *widths: Length) -> None:
    """Fix ``table``'s column widths.

    Both the grid and the cells carry a width: Word prefers the grid, older
    readers the cells, and disagreement between them shows as a wrong layout.
    """
    table.autofit = False
    for index, width in enumerate(widths):
        table.columns[index].width = width
        for cell in table.columns[index].cells:
            cell.width = width


def remove_paragraph(paragraph: Paragraph) -> None:
    """Delete ``paragraph`` from its parent.

    A new table cell arrives with one empty paragraph. Filling the cell appends
    after it, so it has to go or the cell opens with a blank line.
    """
    element = paragraph._p
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def start_hyperlink(paragraph: Paragraph, url: str) -> BaseOxmlElement:
    """Open an external hyperlink at the end of ``paragraph``.

    Returns the ``w:hyperlink`` element; runs become part of the link only once
    they are moved into it with :func:`move_run_into`. Splitting it this way lets
    a link carry nested formatting (bold, code, ...) built with the normal
    python-docx run API.
    """
    relationship_id = paragraph.part.relate_to(url, HYPERLINK_RELATIONSHIP, is_external=True)
    element = OxmlElement("w:hyperlink")
    element.set(qn("r:id"), relationship_id)
    paragraph._p.append(element)
    return element


def move_run_into(run: Run, container: BaseOxmlElement) -> None:
    """Reparent ``run`` into ``container``, e.g. a ``w:hyperlink``.

    lxml moves rather than copies on append, so the run leaves its old position.
    """
    container.append(run._r)


def set_language(style: BaseStyle, lang: str) -> None:
    """Tag ``style`` with a language so Word spell-checks in the right one.

    ``w:lang`` sorts near the end of ``CT_RPr``, after everything python-docx
    writes, so appending is already in schema order.
    """
    properties = style.element.get_or_add_rPr()
    element = OxmlElement("w:lang")
    element.set(qn("w:val"), lang)
    properties.append(element)


# Word's own default z-order (`relativeHeight`) for the first floating shape
# in a document -- confirmed against a real anchored picture in
# data/cv/cv-head.docx. `float_picture` offsets it by the picture's own
# `docPr` id, so two floating pictures in one Document still get distinct
# values without this module having to remember state across calls.
_FLOAT_RELATIVE_HEIGHT_BASE = 251658240

# `CT_Anchor`'s child sequence -- confirmed the same way, against the same
# real picture: `wp:simplePos`, `wp:positionH`, `wp:positionV`, `wp:extent`,
# `wp:effectExtent`, the wrap choice, then `wp:docPr`/`wp:cNvGraphicFramePr`/
# `a:graphic`, appended by `float_picture` itself (see there). `distT`/`distB`/
# `distL`/`distR` (the gap kept between the picture and the text it wraps) and
# `wp:effectExtent`'s all-zero edges are the same values Word itself writes
# for a plain, unshadowed picture.
_ANCHOR_XML = (
    '<wp:anchor {nsdecls} distT="45720" distB="45720" distL="91440" distR="91440" '
    'simplePos="0" relativeHeight="{relative_height}" behindDoc="0" locked="0" '
    'layoutInCell="1" allowOverlap="1">'
    '<wp:simplePos x="0" y="0"/>'
    '<wp:positionH relativeFrom="column"><wp:align>{side}</wp:align></wp:positionH>'
    '<wp:positionV relativeFrom="paragraph"><wp:align>top</wp:align></wp:positionV>'
    '<wp:extent cx="{cx}" cy="{cy}"/>'
    '<wp:effectExtent l="0" t="0" r="0" b="0"/>'
    '<wp:wrapSquare wrapText="bothSides"/>'
    "</wp:anchor>"
)


def float_picture(run: Run, *, side: Literal["left", "right"]) -> None:
    """Turn the inline picture ``run`` just got from ``add_picture`` into a
    floating one, wrapped square and positioned to ``side`` of the column --
    the same shape Word itself writes for a picture whose wrap is changed
    from "in line with text" to "square" and then nudged to a side.

    python-docx has no API for a floating picture (``wp:anchor``) -- only
    ``add_picture``'s inline ``wp:inline`` -- so this rebuilds the wrapper by
    hand (``_ANCHOR_XML``) and moves across the parts that do not change --
    ``wp:docPr``, ``wp:cNvGraphicFramePr`` and ``a:graphic`` (the picture
    itself, ``add_picture`` already built it) -- rather than rebuilding them;
    only the position/wrap/z-order parts, new to a floating picture, are
    freshly written.
    """
    drawing = run._r.find(qn("w:drawing"))
    inline = drawing.find(qn("wp:inline"))
    extent = inline.find(qn("wp:extent"))
    doc_pr = inline.find(qn("wp:docPr"))
    frame_pr = inline.find(qn("wp:cNvGraphicFramePr"))
    graphic = inline.find(qn("a:graphic"))

    relative_height = _FLOAT_RELATIVE_HEIGHT_BASE + int(doc_pr.get("id", "0"))
    anchor = parse_xml(
        _ANCHOR_XML.format(
            nsdecls=nsdecls("wp", "a", "pic", "r"),
            relative_height=relative_height,
            side=side,
            cx=extent.get("cx"),
            cy=extent.get("cy"),
        )
    )
    anchor.append(doc_pr)
    if frame_pr is not None:
        anchor.append(frame_pr)
    anchor.append(graphic)
    drawing.replace(inline, anchor)


def set_picture_crop(
    run: Run, *, left: float = 0.0, top: float = 0.0, right: float = 0.0, bottom: float = 0.0
) -> None:
    """Crop the run's just-added picture the way the source's own crop did.

    Each edge is a fraction (0..1) of the picture, the same shape
    :mod:`cv_generator.docx_import` reads a source crop into. Writes
    ``a:srcRect`` -- Word's own crop record -- rather than pre-cropping the
    bitmap, since that is exactly what a cropped picture in a real ``.docx``
    already carries. It sits between ``a:blip`` and ``a:stretch`` inside
    ``pic:blipFill``, the same fixed sequence python-docx's own
    ``CT_BlipFillProperties.blip`` declares as its successors without
    exposing a way to insert into it. A zero-valued edge is left out rather
    than written as ``="0"``, the same as a source picture that was cropped on
    some edges and not others.
    """
    blip_fill = run._r.find(f".//{qn('pic:blipFill')}")
    if blip_fill is None:
        return
    src_rect = OxmlElement("a:srcRect")
    for edge, value in (("l", left), ("t", top), ("r", right), ("b", bottom)):
        if value:
            src_rect.set(edge, str(round(value * 100000)))
    blip = blip_fill.find(qn("a:blip"))
    blip.addnext(src_rect)
